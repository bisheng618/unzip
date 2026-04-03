import os
import sys
from glob import glob
from typing import List, Optional

from openpyxl import load_workbook


def escape_md(text: Optional[str]) -> str:
    return (str(text) if text is not None else "").replace('|', '\\|').replace('\n', '<br>')


def trim_row(row: List[Optional[object]]) -> List[Optional[object]]:
    i = len(row)
    while i > 0 and (row[i - 1] is None or str(row[i - 1]).strip() == ''):
        i -= 1
    return row[:i]


def rows_to_md(rows: List[List[Optional[object]]]) -> str:
    if not rows:
        return ''
    trimmed = [trim_row(list(r)) for r in rows]
    max_cols = max((len(r) for r in trimmed), default=0)
    if max_cols == 0:
        return ''
    def fmt_row(r: List[Optional[object]]) -> str:
        cells = [escape_md(c) for c in r] + [''] * (max_cols - len(r))
        return '| ' + ' | '.join(cells) + ' |'
    header = fmt_row(trimmed[0])
    sep = '|' + ' | '.join(['---'] * max_cols) + ' |'
    lines = [header, sep]
    for r in trimmed[1:]:
        lines.append(fmt_row(r))
    return '\n'.join(lines) + '\n'


def convert_excel_to_md(excel_path: str, md_path: str) -> None:
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    content = rows_to_md(rows)
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(content)


def convert_directory(dir_path: str) -> None:
    xlsx_files = glob(os.path.join(dir_path, '*.xlsx'))
    if not xlsx_files:
        print(f"No .xlsx files found in: {dir_path}")
        return
    for excel_path in xlsx_files:
        base, _ = os.path.splitext(excel_path)
        md_path = base + '.md'
        convert_excel_to_md(excel_path, md_path)
        print(f"Wrote: {md_path}")


def main():
    default_dir = '/Users/bisheng/TraeProject/tender_generator/excel'
    target_dir = sys.argv[1] if len(sys.argv) > 1 else default_dir
    convert_directory(target_dir)


if __name__ == '__main__':
    main()
