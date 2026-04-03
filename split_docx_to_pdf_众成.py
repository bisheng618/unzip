#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word文档拆分工具 - 完全保留格式版本
直接操作Word文档的XML结构，完全保留原始格式
"""

import os
import subprocess
import shutil
import re
from docx import Document
from docx.oxml import parse_xml
from lxml import etree
import zipfile
import tempfile


class DocxSplitterV2:
    def __init__(self, source_file, output_dir):
        self.source_file = source_file
        self.output_dir = output_dir
        self.doc = Document(source_file)
        
    def get_paragraph_style_level(self, paragraph):
        """获取段落的标题级别"""
        style_name = paragraph.style.name
        if 'Heading' in style_name or '标题' in style_name:
            match = re.search(r'(\d+)', style_name)
            if match:
                return int(match.group(1))
        return None
    
    def find_paragraph_indices(self, start_keyword, end_keyword=None, level=None):
        """查找段落的起始和结束索引"""
        start_idx = None
        end_idx = None
        
        for i, para in enumerate(self.doc.paragraphs):
            if start_keyword in para.text:
                start_idx = i
            elif end_keyword and end_keyword in para.text and start_idx is not None:
                end_idx = i
                break
            elif start_idx is not None and level:
                para_level = self.get_paragraph_style_level(para)
                if para_level and para_level <= level:
                    end_idx = i
                    break
        
        return start_idx, end_idx
    
    def find_level3_sections_by_parent(self, parent_para_idx):
        """找到某个段落下的所有3级标题章节"""
        sections = []
        
        # 获取父段落的级别
        parent_para = self.doc.paragraphs[parent_para_idx]
        parent_level = self.get_paragraph_style_level(parent_para)
        
        current_section = None
        
        for i in range(parent_para_idx + 1, len(self.doc.paragraphs)):
            para = self.doc.paragraphs[i]
            level = self.get_paragraph_style_level(para)
            
            # 如果遇到同级或更高级别的标题，停止
            if level and level <= parent_level:
                if current_section:
                    current_section['end_idx'] = i
                    sections.append(current_section)
                break
            
            # 找到3级标题
            if level == 3:
                # 保存上一个section
                if current_section:
                    current_section['end_idx'] = i
                    sections.append(current_section)
                
                # 开始新section
                current_section = {
                    'title': para.text.strip(),
                    'start_idx': i,
                    'end_idx': None
                }
        
        # 添加最后一个section
        if current_section and not current_section['end_idx']:
            current_section['end_idx'] = len(self.doc.paragraphs)
            sections.append(current_section)
        
        return sections
    
    def extract_docx_by_paragraph_range(self, start_para_idx, end_para_idx, output_path):
        """通过复制整个docx文件并删除不需要的段落来提取内容，完全保留格式"""
        # 创建临时目录
        with tempfile.TemporaryDirectory() as temp_dir:
            # 复制原始文档到临时位置
            temp_docx = os.path.join(temp_dir, 'temp.docx')
            shutil.copy2(self.source_file, temp_docx)
            
            # 打开复制的文档
            doc = Document(temp_docx)
            
            # 获取所有段落和它们的父元素
            body = doc.element.body
            
            # 收集要删除的元素
            elements_to_remove = []
            
            # Word命名空间
            w_namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            
            # 先获取所有段落的样式信息
            para_styles = []
            for para in doc.paragraphs:
                para_styles.append(para.style.name.lower())
            
            # 遍历body的所有子元素
            para_count = 0
            for element in list(body):
                # 检查是否是段落
                if element.tag.endswith('p'):
                    should_remove = False
                    
                    # 检查是否在范围外
                    if para_count < start_para_idx or (end_para_idx and para_count >= end_para_idx):
                        should_remove = True
                    else:
                        # 检查段落样式是否包含'toc'（目录样式）
                        if para_count < len(para_styles):
                            style_name = para_styles[para_count]
                            if 'toc' in style_name:
                                should_remove = True
                        
                        # 检查是否包含TOC字段指令
                        if not should_remove:
                            for child in element.iter():
                                # 检查instrText元素
                                if child.tag == f'{w_namespace}instrText':
                                    if child.text and 'TOC' in child.text:
                                        should_remove = True
                                        break
                        
                        # 如果段落文本为"目录"，也删除
                        if not should_remove:
                            para_text = ''
                            for child in element.iter():
                                if child.tag == f'{w_namespace}t':
                                    if child.text:
                                        para_text += child.text
                            
                            if para_text.strip() in ['目录', 'Table of Contents', 'TOC']:
                                should_remove = True
                    
                    if should_remove:
                        elements_to_remove.append(element)
                    
                    para_count += 1
                    
                # 如果是表格，检查它是否在范围内
                elif element.tag.endswith('tbl'):
                    # 表格比较复杂，我们需要检查它前后的段落位置
                    # 为了简单起见，如果表格在段落范围内就保留
                    if para_count < start_para_idx or (end_para_idx and para_count >= end_para_idx):
                        elements_to_remove.append(element)
            
            # 删除不需要的元素
            for element in elements_to_remove:
                body.remove(element)
            
            # 保存修改后的文档
            doc.save(output_path)
    
    def convert_docx_to_pdf(self, docx_path, pdf_path):
        """使用LibreOffice将DOCX转换为PDF"""
        try:
            cmd = [
                '/Applications/LibreOffice.app/Contents/MacOS/soffice',
                '--headless',
                '--convert-to', 'pdf',
                '--outdir', os.path.dirname(pdf_path),
                docx_path
            ]
            
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
            
            # LibreOffice会生成与输入文件同名的PDF
            generated_pdf = docx_path.replace('.docx', '.pdf')
            
            if os.path.exists(generated_pdf):
                if generated_pdf != pdf_path:
                    shutil.move(generated_pdf, pdf_path)
                print(f"✓ 已生成: {os.path.basename(pdf_path)}")
                return True
            else:
                print(f"✗ 转换失败: {os.path.basename(pdf_path)}")
                return False
                
        except subprocess.TimeoutExpired:
            print(f"✗ 转换超时: {os.path.basename(pdf_path)}")
            return False
        except Exception as e:
            print(f"✗ 转换出错 {os.path.basename(pdf_path)}: {e}")
            return False
    
    def save_section_as_pdf(self, start_idx, end_idx, output_path):
        """保存段落范围为PDF"""
        temp_docx = output_path.replace('.pdf', '_temp.docx')
        
        # 提取内容到临时docx
        self.extract_docx_by_paragraph_range(start_idx, end_idx, temp_docx)
        
        # 转换为PDF
        success = self.convert_docx_to_pdf(temp_docx, output_path)
        
        # 删除临时文件
        if os.path.exists(temp_docx):
            os.remove(temp_docx)
        
        return success
    
    def split_performance_files(self):
        """拆分"业绩文件"下的3个小节"""
        print("\n1. 拆分业绩文件...")
        output_folder = os.path.join(self.output_dir, "业绩文件")
        os.makedirs(output_folder, exist_ok=True)
        
        parent_idx = None
        for i, para in enumerate(self.doc.paragraphs):
            if "业绩文件" in para.text:
                parent_idx = i
                break
        
        if parent_idx is None:
            print("  未找到'业绩文件'")
            return
        
        sections = self.find_level3_sections_by_parent(parent_idx)
        print(f"  找到 {len(sections)} 个小节")
        
        for idx, section in enumerate(sections, 1):
            title = section['title']
            safe_title = re.sub(r'[\\/:*?"<>|]', '_', title)
            output_path = os.path.join(output_folder, f"{idx:02d}_{safe_title}.pdf")
            self.save_section_as_pdf(section['start_idx'], section['end_idx'], output_path)
    
    def split_management_summary(self):
        """拆分"拟委任的主要人员汇总表"部分"""
        print("\n2. 拆分拟委任的主要人员汇总表...")
        
        parent_idx = None
        for i, para in enumerate(self.doc.paragraphs):
            if "拟委任的主要人员汇总表" in para.text:
                parent_idx = i
                break
        
        if parent_idx is None:
            print("  未找到'拟委任的主要人员汇总表'")
            return
        
        # 找到下一个同级或更高级别的标题
        end_idx = None
        parent_level = self.get_paragraph_style_level(self.doc.paragraphs[parent_idx])
        for i in range(parent_idx + 1, len(self.doc.paragraphs)):
            level = self.get_paragraph_style_level(self.doc.paragraphs[i])
            if level and level <= parent_level:
                end_idx = i
                break
        
        output_path = os.path.join(self.output_dir, "拟委任的主要人员汇总表.pdf")
        self.save_section_as_pdf(parent_idx, end_idx, output_path)
    
    def split_resume_table(self):
        """拆分"表2主要人员简历表"部分"""
        print("\n3. 拆分表2主要人员简历表...")
        
        parent_idx = None
        for i, para in enumerate(self.doc.paragraphs):
            if "表2主要人员简历表" in para.text or "表2" in para.text and "主要人员简历" in para.text:
                parent_idx = i
                break
        
        if parent_idx is None:
            print("  未找到'表2主要人员简历表'")
            return
        
        # 找到下一个同级或更高级别的标题
        end_idx = None
        parent_level = self.get_paragraph_style_level(self.doc.paragraphs[parent_idx])
        for i in range(parent_idx + 1, len(self.doc.paragraphs)):
            level = self.get_paragraph_style_level(self.doc.paragraphs[i])
            if level and level <= parent_level:
                end_idx = i
                break
        
        output_path = os.path.join(self.output_dir, "表2主要人员简历表.pdf")
        self.save_section_as_pdf(parent_idx, end_idx, output_path)
    
    def split_technical_sections(self):
        """拆分"技术方案"到"后续服务承诺及措施"为1个文件"""
        print("\n4. 拆分技术方案到后续服务承诺及措施...")
        
        start_idx = None
        end_idx = None
        
        for i, para in enumerate(self.doc.paragraphs):
            text = para.text.strip()
            level = self.get_paragraph_style_level(para)
            
            # 查找"技术方案"标题（level 1）
            if start_idx is None and "技术方案" in text and level == 1:
                start_idx = i
            
            # 查找"后续服务承诺及措施"之后的下一个同级标题
            if start_idx and "后续服务承诺及措施" in text and level == 1:
                # 找到后续服务承诺及措施之后的下一个level 1标题
                for j in range(i + 1, len(self.doc.paragraphs)):
                    next_level = self.get_paragraph_style_level(self.doc.paragraphs[j])
                    if next_level and next_level <= 1:
                        end_idx = j
                        break
                break
        
        if start_idx:
            output_path = os.path.join(self.output_dir, "技术方案到后续服务承诺及措施.pdf")
            self.save_section_as_pdf(start_idx, end_idx, output_path)
        else:
            print("  未找到技术方案章节")
    
    def split_audit_reports(self):
        """拆分2022、2023、2024年审计报告为3个文件"""
        print("\n5. 拆分审计报告...")
        
        years = ["2022", "2023", "2024"]
        
        for year in years:
            start_idx = None
            end_idx = None
            
            for i, para in enumerate(self.doc.paragraphs):
                text = para.text.strip()
                
                # 查找该年份的审计报告
                if f"{year}年审计报告" in text or f"{year}审计报告" in text:
                    start_idx = i
                    
                    # 找到下一个同级或更高级别的标题
                    parent_level = self.get_paragraph_style_level(self.doc.paragraphs[i])
                    for j in range(i + 1, len(self.doc.paragraphs)):
                        level = self.get_paragraph_style_level(self.doc.paragraphs[j])
                        # 检查是否是下一年的审计报告或其他同级标题
                        next_text = self.doc.paragraphs[j].text.strip()
                        if (level and level <= parent_level) or "审计报告" in next_text:
                            end_idx = j
                            break
                    break
            
            if start_idx:
                output_path = os.path.join(self.output_dir, f"{year}年审计报告.pdf")
                self.save_section_as_pdf(start_idx, end_idx, output_path)
            else:
                print(f"  未找到{year}年审计报告")
    
    def run(self):
        """执行所有拆分任务"""
        print("=" * 80)
        print("开始拆分文档（完全保留格式）...")
        print("=" * 80)
        
        self.split_performance_files()
        self.split_management_summary()
        self.split_resume_table()
        self.split_technical_sections()
        self.split_audit_reports()
        
        print("\n" + "=" * 80)
        print("所有拆分任务完成！")
        print("=" * 80)


def main():
    source_file = "/Users/bisheng/Downloads/zc/包2-技术文件.docx"
    output_dir = "/Users/bisheng/Downloads/zc/"
    
    if not os.path.exists(source_file):
        print(f"错误: 源文件不存在: {source_file}")
        return
    
    if not os.path.exists('/Applications/LibreOffice.app/Contents/MacOS/soffice'):
        print("错误: 未找到LibreOffice，请先安装LibreOffice")
        print("下载地址: https://www.libreoffice.org/download/download/")
        return
    
    splitter = DocxSplitterV2(source_file, output_dir)
    splitter.run()


if __name__ == "__main__":
    main()
