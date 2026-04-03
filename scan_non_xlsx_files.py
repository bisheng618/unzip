#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
扫描指定目录下所有非xlsx结尾的文件,并将xls文件转换为xlsx格式
"""

import os
import logging
from pathlib import Path
import pandas as pd

# 配置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('non_xlsx_files.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)

logger = logging.getLogger(__name__)


def convert_xls_to_xlsx(xls_file_path):
    """
    将xls文件转换为xlsx格式,保留原始格式
    
    Args:
        xls_file_path: xls文件的完整路径
    
    Returns:
        bool: 转换是否成功
    """
    try:
        import xlrd
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        from openpyxl.utils import get_column_letter
        
        # 生成新的xlsx文件路径
        xlsx_file_path = xls_file_path.rsplit('.', 1)[0] + '.xlsx'
        
        logger.info(f"开始转换: {xls_file_path}")
        
        # 打开xls文件
        xls_book = xlrd.open_workbook(xls_file_path, formatting_info=True)
        
        # 创建新的xlsx工作簿
        xlsx_book = Workbook()
        xlsx_book.remove(xlsx_book.active)  # 删除默认创建的sheet
        
        # 遍历所有sheet
        for sheet_index in range(xls_book.nsheets):
            xls_sheet = xls_book.sheet_by_index(sheet_index)
            sheet_name = xls_sheet.name
            
            # 创建新sheet
            xlsx_sheet = xlsx_book.create_sheet(title=sheet_name)
            
            # 复制数据和格式
            for row_idx in range(xls_sheet.nrows):
                for col_idx in range(xls_sheet.ncols):
                    xls_cell = xls_sheet.cell(row_idx, col_idx)
                    xlsx_cell = xlsx_sheet.cell(row=row_idx + 1, column=col_idx + 1)
                    
                    # 复制单元格值
                    cell_value = xls_cell.value
                    
                    # 处理日期类型
                    if xls_cell.ctype == 3:  # 日期类型
                        try:
                            from datetime import datetime
                            cell_value = xlrd.xldate_as_datetime(xls_cell.value, xls_book.datemode)
                        except:
                            pass
                    
                    xlsx_cell.value = cell_value
                    
                    # 复制格式
                    try:
                        xls_xf = xls_book.xf_list[xls_cell.xf_index]
                        
                        # 获取字体信息
                        font_index = xls_xf.font_index
                        xls_font = xls_book.font_list[font_index]
                        
                        # 获取字体颜色
                        font_color = None
                        if xls_font.colour_index != 32767:  # 32767表示默认颜色
                            try:
                                color_map = xls_book.colour_map.get(xls_font.colour_index)
                                if color_map:
                                    # 转换RGB颜色
                                    font_color = '%02x%02x%02x' % color_map
                            except:
                                pass
                        
                        # 设置字体
                        xlsx_cell.font = Font(
                            name=xls_font.name,
                            size=xls_font.height / 20,  # xlrd中字体大小单位是twips
                            bold=xls_font.bold,
                            italic=xls_font.italic,
                            underline='single' if xls_font.underline_type else 'none',
                            color=font_color
                        )
                        
                        # 获取对齐信息
                        alignment_dict = {
                            0: 'general',
                            1: 'left',
                            2: 'center',
                            3: 'right',
                            4: 'fill',
                            5: 'justify',
                            6: 'centerContinuous',
                            7: 'distributed'
                        }
                        
                        vertical_dict = {
                            0: 'top',
                            1: 'center',
                            2: 'bottom',
                            3: 'justify',
                            4: 'distributed'
                        }
                        
                        horizontal = alignment_dict.get(xls_xf.alignment.hor_align, 'general')
                        vertical = vertical_dict.get(xls_xf.alignment.vert_align, 'bottom')
                        
                        xlsx_cell.alignment = Alignment(
                            horizontal=horizontal,
                            vertical=vertical,
                            wrap_text=bool(xls_xf.alignment.text_wrapped)
                        )
                        
                        # 复制背景色
                        if xls_xf.background.pattern_colour_index != 64:  # 64表示无背景
                            try:
                                bg_color = xls_book.colour_map.get(xls_xf.background.pattern_colour_index)
                                if bg_color:
                                    bg_color_hex = '%02x%02x%02x' % bg_color
                                    xlsx_cell.fill = PatternFill(
                                        start_color=bg_color_hex,
                                        end_color=bg_color_hex,
                                        fill_type='solid'
                                    )
                            except:
                                pass
                        
                        # 复制边框
                        border_style_map = {
                            0: None,  # 无边框
                            1: 'thin',
                            2: 'medium',
                            3: 'dashed',
                            4: 'dotted',
                            5: 'thick',
                            6: 'double',
                            7: 'hair',
                            8: 'mediumDashed',
                            9: 'dashDot',
                            10: 'mediumDashDot',
                            11: 'dashDotDot',
                            12: 'mediumDashDotDot',
                            13: 'slantDashDot'
                        }
                        
                        def get_border_side(line_style, color_index):
                            """获取边框样式"""
                            if line_style == 0:
                                return Side(style=None)
                            
                            style = border_style_map.get(line_style, 'thin')
                            border_color = None
                            
                            if color_index != 64 and color_index != 32767:
                                try:
                                    color_rgb = xls_book.colour_map.get(color_index)
                                    if color_rgb:
                                        border_color = '%02x%02x%02x' % color_rgb
                                except:
                                    pass
                            
                            return Side(style=style, color=border_color)
                        
                        # 设置边框
                        xlsx_cell.border = Border(
                            left=get_border_side(xls_xf.border.left_line_style, xls_xf.border.left_colour_index),
                            right=get_border_side(xls_xf.border.right_line_style, xls_xf.border.right_colour_index),
                            top=get_border_side(xls_xf.border.top_line_style, xls_xf.border.top_colour_index),
                            bottom=get_border_side(xls_xf.border.bottom_line_style, xls_xf.border.bottom_colour_index)
                        )
                        
                    except Exception as e:
                        # 如果格式复制失败,继续处理下一个单元格
                        pass
            
            # 复制列宽
            for col_idx in range(xls_sheet.ncols):
                try:
                    col_width = xls_sheet.colinfo_map.get(col_idx)
                    if col_width:
                        # xlrd中列宽单位是256分之一字符宽度
                        xlsx_sheet.column_dimensions[get_column_letter(col_idx + 1)].width = col_width.width / 256
                except:
                    pass
            
            # 复制行高
            for row_idx in range(xls_sheet.nrows):
                try:
                    row_info = xls_sheet.rowinfo_map.get(row_idx)
                    if row_info:
                        # xlrd中行高单位是twips (1/20 point)
                        xlsx_sheet.row_dimensions[row_idx + 1].height = row_info.height / 20
                except:
                    pass
            
            # 复制合并单元格
            for crange in xls_sheet.merged_cells:
                rlo, rhi, clo, chi = crange
                xlsx_sheet.merge_cells(
                    start_row=rlo + 1,
                    start_column=clo + 1,
                    end_row=rhi,
                    end_column=chi
                )
        
        # 保存xlsx文件
        xlsx_book.save(xlsx_file_path)
        
        logger.info(f"转换成功: {xlsx_file_path}")
        
        # 删除原xls文件
        os.remove(xls_file_path)
        logger.info(f"已删除原文件: {xls_file_path}")
        
        return True
        
    except Exception as e:
        logger.error(f"转换失败 {xls_file_path}: {str(e)}")
        import traceback
        logger.error(traceback.format_exc())
        return False


def scan_and_convert_files(target_dir):
    """
    扫描目标目录下所有非xlsx结尾的文件,并转换xls文件
    
    Args:
        target_dir: 目标目录路径
    """
    target_path = Path(target_dir)
    
    if not target_path.exists():
        logger.error(f"目标目录不存在: {target_dir}")
        return
    
    if not target_path.is_dir():
        logger.error(f"目标路径不是目录: {target_dir}")
        return
    
    logger.info(f"开始扫描目录: {target_dir}")
    logger.info("=" * 80)
    
    xls_files = []
    non_xlsx_count = 0
    total_count = 0
    
    # 遍历目录下所有文件
    for root, dirs, files in os.walk(target_path):
        for file in files:
            total_count += 1
            file_path = os.path.join(root, file)
            
            # 检查文件是否以.xlsx结尾
            if not file.lower().endswith('.xlsx'):
                non_xlsx_count += 1
                logger.info(f"非xlsx文件: {file_path}")
                
                # 如果是xls文件,添加到转换列表
                if file.lower().endswith('.xls'):
                    xls_files.append(file_path)
    
    logger.info("=" * 80)
    logger.info(f"扫描完成! 总文件数: {total_count}, 非xlsx文件数: {non_xlsx_count}")
    logger.info(f"发现 {len(xls_files)} 个xls文件需要转换")
    
    # 转换xls文件
    if xls_files:
        logger.info("=" * 80)
        logger.info("开始转换xls文件...")
        success_count = 0
        fail_count = 0
        
        for xls_file in xls_files:
            if convert_xls_to_xlsx(xls_file):
                success_count += 1
            else:
                fail_count += 1
        
        logger.info("=" * 80)
        logger.info(f"转换完成! 成功: {success_count}, 失败: {fail_count}")


def main():
    # 目标目录
    target_directory = "/Users/bisheng/Downloads/12171/工程量清单_汇总"
    
    scan_and_convert_files(target_directory)


if __name__ == "__main__":
    main()
