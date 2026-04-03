#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
基于项目列表的文件匹配与复制程序

功能:
1. 读取 project.txt 中的项目名称列表
2. 读取 Excel 文件中的"文件名"和"文件路径"列
3. 过滤 Excel 数据,仅处理文件名包含 project.txt 中任意项目名称的行
4. 在指定目录下查找这些项目对应的包含"完成报价"的文件
5. 将匹配的文件复制到 Excel 中"文件路径"列指定位置下的"上传版"路径中
"""

import os
import shutil
import logging
from pathlib import Path
from datetime import datetime
import openpyxl
from typing import Dict, List, Set


class ProjectFileMatcher:
    """基于项目列表的文件匹配与复制类"""
    
    def __init__(self, excel_path: str, search_dir: str, project_txt_path: str):
        """
        初始化
        
        Args:
            excel_path: Excel 文件路径
            search_dir: 搜索目录路径
            project_txt_path: project.txt 文件路径
        """
        self.excel_path = excel_path
        self.search_dir = search_dir
        self.project_txt_path = project_txt_path
        self.setup_logging()
        
        # 统计信息
        self.stats = {
            'total_projects_in_txt': 0,
            'total_excel_rows': 0,
            'filtered_rows': 0,
            'matched_files': 0,
            'copied_files': 0,
            'failed_copies': 0,
            'no_match': 0,
            'skipped_existing': 0
        }
    
    def setup_logging(self):
        """设置日志"""
        log_filename = f'project_copy_log_{datetime.now().strftime("%Y%m%d_%H%M%S")}.txt'
        
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler(log_filename, encoding='utf-8'),
                logging.StreamHandler()
            ]
        )
        self.logger = logging.getLogger(__name__)
        self.logger.info(f"日志文件: {log_filename}")
    
    def read_project_list(self) -> Set[str]:
        """读取 project.txt 中的项目名称"""
        if not os.path.exists(self.project_txt_path):
            self.logger.error(f"找不到项目列表文件: {self.project_txt_path}")
            return set()
        
        with open(self.project_txt_path, 'r', encoding='utf-8') as f:
            projects = {line.strip() for line in f if line.strip()}
        
        self.stats['total_projects_in_txt'] = len(projects)
        self.logger.info(f"已加载 {len(projects)} 个项目关键词")
        return projects

    def read_and_filter_excel(self, project_names: Set[str]) -> Dict[str, Dict]:
        """读取并按项目名过滤 Excel 数据"""
        self.logger.info(f"开始读取 Excel 文件: {self.excel_path}")
        
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel文件不存在: {self.excel_path}")
        
        try:
            workbook = openpyxl.load_workbook(self.excel_path)
            sheet = workbook.active
            
            # 查找"文件名"和"文件路径"列
            filename_col = None
            filepath_col = None
            header_row = None
            
            for row_idx, row in enumerate(sheet.iter_rows(max_row=10), 1):
                for col_idx, cell in enumerate(row, 1):
                    if cell.value:
                        val = str(cell.value).strip()
                        if '文件名' in val:
                            filename_col = col_idx
                            header_row = row_idx
                        elif '文件路径' in val:
                            filepath_col = col_idx
                if filename_col and filepath_col:
                    break
            
            if not filename_col or not filepath_col:
                raise ValueError("未找到'文件名'或'文件路径'列")
            
            data = {}
            for row in sheet.iter_rows(min_row=header_row + 1):
                filename_cell = row[filename_col - 1]
                filepath_cell = row[filepath_col - 1]
                
                if not filename_cell.value:
                    continue
                
                original_filename = str(filename_cell.value).strip()
                self.stats['total_excel_rows'] += 1
                
                # 过滤: 检查文件名是否包含 project.txt 中的任意项目名
                matched_project = None
                for proj in project_names:
                    if proj in original_filename:
                        matched_project = proj
                        break
                
                if not matched_project:
                    continue
                
                # 去除扩展名作为匹配键
                base_name = os.path.splitext(original_filename)[0]
                file_path = str(filepath_cell.value).strip() if filepath_cell.value else ""
                
                data[base_name] = {
                    'original_filename': original_filename,
                    'file_path': file_path,
                    'matched_project': matched_project,
                    'matching_files_on_disk': []
                }
                self.stats['filtered_rows'] += 1
            
            workbook.close()
            self.logger.info(f"Excel 总行数: {self.stats['total_excel_rows']}, 过滤后匹配项目数的行数: {self.stats['filtered_rows']}")
            return data
            
        except Exception as e:
            self.logger.error(f"读取或过滤 Excel 失败: {e}")
            raise

    def search_files_on_disk(self, data: Dict[str, Dict]):
        """在磁盘上搜索匹配的文件"""
        self.logger.info(f"开始在目录中搜索文件: {self.search_dir}")
        
        if not os.path.exists(self.search_dir):
            raise FileNotFoundError(f"搜索目录不存在: {self.search_dir}")
        
        search_path = Path(self.search_dir)
        all_files = list(search_path.rglob('*'))
        self.logger.info(f"目录中共发现 {len(all_files)} 个项")
        
        for file_item in all_files:
            if not file_item.is_file():
                continue
            
            filename = file_item.name
            
            # 条件: 必须包含 "完成报价"
            if "完成报价" not in filename:
                continue
            
            # 检查是否匹配 Excel 中的某个条目
            for base_name, info in data.items():
                if base_name in filename:
                    info['matching_files_on_disk'].append(str(file_item))
                    self.stats['matched_files'] += 1
                    self.logger.info(f"磁盘匹配成功: [{base_name}] -> {filename}")
        
        # 统计没有在磁盘找到对应文件的条目
        for base_name, info in data.items():
            if not info['matching_files_on_disk']:
                self.stats['no_match'] += 1
                self.logger.warning(f"未在磁盘找到匹配文件: {base_name} (项目: {info['matched_project']})")

    def copy_matched_files(self, data: Dict[str, Dict]):
        """执行复制操作"""
        self.logger.info("开始执行文件拷贝")
        
        for base_name, info in data.items():
            if not info['matching_files_on_disk']:
                continue
            
            if not info['file_path']:
                self.logger.warning(f"Excel 中路径为空,跳过: {base_name}")
                continue
            
            # 目标目录: Excel路径 + "上传版"
            target_dir = os.path.join(info['file_path'], "上传版")
            
            # 检查目标目录是否已经有报价文件
            if os.path.exists(target_dir):
                existing = False
                for f in os.listdir(target_dir):
                    if "完成报价" in f and (f.endswith(".xlsx") or f.endswith(".xls")):
                        existing = True
                        break
                if existing:
                    self.logger.info(f"目标目录已存在报价文件,跳过: {base_name}")
                    self.stats['skipped_existing'] += 1
                    continue
            
            try:
                os.makedirs(target_dir, exist_ok=True)
            except Exception as e:
                self.logger.error(f"创建目录失败 {target_dir}: {e}")
                self.stats['failed_copies'] += len(info['matching_files_on_disk'])
                continue
            
            for src in info['matching_files_on_disk']:
                try:
                    dest = os.path.join(target_dir, os.path.basename(src))
                    shutil.copy2(src, dest)
                    self.stats['copied_files'] += 1
                    self.logger.info(f"复制成功: {os.path.basename(src)} -> {target_dir}")
                except Exception as e:
                    self.logger.error(f"复制失败 {src}: {e}")
                    self.stats['failed_copies'] += 1

    def run(self):
        """执行全流程"""
        try:
            self.logger.info("="*30)
            self.logger.info("项目过滤拷贝程序启动")
            self.logger.info(f"项目列表: {self.project_txt_path}")
            self.logger.info(f"Excel 路径: {self.excel_path}")
            self.logger.info(f"搜索目录: {self.search_dir}")
            self.logger.info("="*30)
            
            # 1. 加载项目列表
            projects = self.read_project_list()
            if not projects:
                self.logger.warning("项目列表为空,程序退出")
                return
            
            # 2. 读取并过滤 Excel
            data = self.read_and_filter_excel(projects)
            
            # 3. 磁盘搜索
            self.search_files_on_disk(data)
            
            # 4. 拷贝文件
            self.copy_matched_files(data)
            
            # 5. 打印总结
            self.print_summary()
            
        except Exception as e:
            self.logger.error(f"程序运行异常: {e}", exc_info=True)

    def print_summary(self):
        """摘要"""
        self.logger.info("\n" + "="*60)
        self.logger.info("处理完成 - 统计摘要")
        self.logger.info("="*60)
        self.logger.info(f"project.txt 项目数: {self.stats['total_projects_in_txt']}")
        self.logger.info(f"Excel 总行数: {self.stats['total_excel_rows']}")
        self.logger.info(f"匹配项目的 Excel 行数: {self.stats['filtered_rows']}")
        self.logger.info(f"磁盘搜索到的匹配文件数: {self.stats['matched_files']}")
        self.logger.info(f"磁盘未匹配的项目行数: {self.stats['no_match']}")
        self.logger.info(f"已存在跳过的文件数: {self.stats['skipped_existing']}")
        self.logger.info(f"成功拷贝的文件数: {self.stats['copied_files']}")
        self.logger.info(f"拷贝失败数: {self.stats['failed_copies']}")
        self.logger.info("="*60)


def main():
    # 路径配置
    excel_path = "/Volumes/zhuxl 共享给我/1、25年年底投标汇总8888/国网安徽电力2026年第一次服务类区域联合授权竞争性谈判采购/工作量清单_文件汇总.xlsx"
    search_dir = "/Volumes/zhuxl 共享给我/1、25年年底投标汇总8888/国网安徽电力2026年第一次服务类区域联合授权竞争性谈判采购/工程量清单_汇总"
    project_txt_path = "project.txt" # 默认在当前目录下
    
    # 实例化并运行
    matcher = ProjectFileMatcher(excel_path, search_dir, project_txt_path)
    matcher.run()


if __name__ == "__main__":
    main()
