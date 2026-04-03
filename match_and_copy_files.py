#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文件匹配与复制程序

功能:
1. 读取Excel文件中的"文件名"列,去除扩展名
2. 在指定目录下查找包含这些文件名的文件
3. 如果文件名包含"完成报价",则复制到Excel中"文件路径"列指定的位置+"上传版"路径下
"""

import os
import shutil
import logging
from pathlib import Path
from datetime import datetime
import openpyxl
from typing import Dict, List, Tuple


class FileMatcherCopier:
    """文件匹配与复制类"""
    
    def __init__(self, excel_path: str, search_dir: str):
        """
        初始化
        
        Args:
            excel_path: Excel文件路径
            search_dir: 搜索目录路径
        """
        self.excel_path = excel_path
        self.search_dir = search_dir
        self.setup_logging()
        
        # 统计信息
        self.stats = {
            'total_rows': 0,
            'matched_files': 0,
            'copied_files': 0,
            'failed_copies': 0,
            'no_match': 0,
            'skipped_existing': 0
        }
    
    def setup_logging(self):
        """设置日志"""
        log_filename = f'file_copy_log_{datetime.now().strftime("%Y%m%d_%H%M%S")}.txt'
        
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler(log_filename, encoding='utf-8'),
                logging.StreamHandler()
            ]
        )
        self.logger = logging.getLogger(__name__)
        self.logger.info(f"日志文件: {log_filename}")
    
    def read_excel_data(self) -> Dict[str, Dict]:
        """
        读取Excel文件数据
        
        Returns:
            字典,key为文件名(无扩展名), value为包含原始文件名和文件路径的字典
        """
        self.logger.info(f"开始读取Excel文件: {self.excel_path}")
        
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel文件不存在: {self.excel_path}")
        
        try:
            workbook = openpyxl.load_workbook(self.excel_path)
            sheet = workbook.active
            
            # 查找"文件名"和"文件路径"列的索引
            header_row = None
            filename_col = None
            filepath_col = None
            
            for row_idx, row in enumerate(sheet.iter_rows(max_row=10), 1):
                for col_idx, cell in enumerate(row, 1):
                    if cell.value:
                        cell_value = str(cell.value).strip()
                        if '文件名' in cell_value:
                            filename_col = col_idx
                            header_row = row_idx
                        elif '文件路径' in cell_value:
                            filepath_col = col_idx
                
                if filename_col and filepath_col:
                    break
            
            if not filename_col or not filepath_col:
                raise ValueError("未找到'文件名'或'文件路径'列")
            
            self.logger.info(f"找到列: 文件名(列{filename_col}), 文件路径(列{filepath_col}), 表头行{header_row}")
            
            # 读取数据
            data = {}
            for row in sheet.iter_rows(min_row=header_row + 1):
                filename_cell = row[filename_col - 1]
                filepath_cell = row[filepath_col - 1]
                
                if filename_cell.value:
                    original_filename = str(filename_cell.value).strip()
                    # 去除扩展名
                    base_filename = os.path.splitext(original_filename)[0]
                    
                    # 获取文件路径
                    file_path = str(filepath_cell.value).strip() if filepath_cell.value else ""
                    
                    data[base_filename] = {
                        'original_filename': original_filename,
                        'file_path': file_path,
                        'matched_files': []
                    }
                    self.stats['total_rows'] += 1
            
            workbook.close()
            self.logger.info(f"成功读取 {self.stats['total_rows']} 行数据")
            return data
            
        except Exception as e:
            self.logger.error(f"读取Excel文件失败: {e}")
            raise
    
    def search_matching_files(self, data: Dict[str, Dict]) -> Dict[str, Dict]:
        """
        搜索匹配的文件
        
        Args:
            data: Excel数据字典
            
        Returns:
            更新后的数据字典,包含匹配的文件列表
        """
        self.logger.info(f"开始在目录中搜索匹配文件: {self.search_dir}")
        
        if not os.path.exists(self.search_dir):
            raise FileNotFoundError(f"搜索目录不存在: {self.search_dir}")
        
        # 遍历搜索目录
        search_path = Path(self.search_dir)
        all_files = list(search_path.rglob('*'))
        self.logger.info(f"搜索目录中共有 {len(all_files)} 个文件/文件夹")
        
        for file_path in all_files:
            if not file_path.is_file():
                continue
            
            filename = file_path.name
            
            # 检查文件名是否包含"完成报价"
            if "完成报价" not in filename:
                continue
            
            # 检查是否匹配Excel中的基础文件名
            for base_filename, info in data.items():
                if base_filename in filename:
                    info['matched_files'].append(str(file_path))
                    self.stats['matched_files'] += 1
                    self.logger.info(f"匹配: [{base_filename}] -> {filename}")
        
        # 统计未匹配的
        for base_filename, info in data.items():
            if not info['matched_files']:
                self.stats['no_match'] += 1
                self.logger.warning(f"未找到匹配文件: {base_filename}")
        
        return data
    
    def copy_files(self, data: Dict[str, Dict]):
        """
        复制文件到目标位置
        
        Args:
            data: 包含匹配文件的数据字典
        """
        self.logger.info("开始复制文件")
        
        for base_filename, info in data.items():
            if not info['matched_files']:
                continue
            
            # 构建目标路径: 文件路径 + "上传版"
            if not info['file_path']:
                self.logger.warning(f"文件路径为空,跳过: {base_filename}")
                continue
            
            target_dir = os.path.join(info['file_path'], "上传版")
            
            # [新增] 检查目标目录是否已存在包含“完成报价”的Excel文件
            if os.path.exists(target_dir):
                existing_quote_file = False
                for f in os.listdir(target_dir):
                    if "完成报价" in f and (f.endswith(".xlsx") or f.endswith(".xls")):
                        existing_quote_file = True
                        break
                
                if existing_quote_file:
                    self.logger.info(f"目标文件夹已存在报价文件,跳过拷贝: {base_filename}")
                    self.stats['skipped_existing'] += 1
                    continue

            # 创建目标目录
            try:
                os.makedirs(target_dir, exist_ok=True)
                self.logger.info(f"目标目录: {target_dir}")
            except Exception as e:
                self.logger.error(f"创建目标目录失败 [{target_dir}]: {e}")
                self.stats['failed_copies'] += len(info['matched_files'])
                continue
            
            # 复制所有匹配的文件
            for source_file in info['matched_files']:
                try:
                    filename = os.path.basename(source_file)
                    target_file = os.path.join(target_dir, filename)
                    
                    shutil.copy2(source_file, target_file)
                    self.stats['copied_files'] += 1
                    self.logger.info(f"复制成功: {filename} -> {target_dir}")
                    
                except Exception as e:
                    self.stats['failed_copies'] += 1
                    self.logger.error(f"复制失败 [{source_file}]: {e}")
    
    def print_summary(self):
        """打印统计摘要"""
        self.logger.info("\n" + "="*60)
        self.logger.info("处理完成 - 统计摘要")
        self.logger.info("="*60)
        self.logger.info(f"Excel总行数: {self.stats['total_rows']}")
        self.logger.info(f"找到匹配文件数: {self.stats['matched_files']}")
        self.logger.info(f"未找到匹配: {self.stats['no_match']}")
        self.logger.info(f"由于已存在且跳过的任务数: {self.stats['skipped_existing']}")
        self.logger.info(f"成功复制文件数: {self.stats['copied_files']}")
        self.logger.info(f"复制失败数: {self.stats['failed_copies']}")
        self.logger.info("="*60)
    
    def run(self):
        """运行主流程"""
        try:
            self.logger.info("程序启动")
            self.logger.info(f"Excel文件: {self.excel_path}")
            self.logger.info(f"搜索目录: {self.search_dir}")
            
            # 1. 读取Excel数据
            data = self.read_excel_data()
            
            # 2. 搜索匹配文件
            data = self.search_matching_files(data)
            
            # 3. 复制文件
            self.copy_files(data)
            
            # 4. 打印摘要
            self.print_summary()
            
            self.logger.info("程序执行完成")
            
        except Exception as e:
            self.logger.error(f"程序执行失败: {e}", exc_info=True)
            raise


def main():
    """主函数"""
    # 配置路径
    excel_path = "/Volumes/zhuxl 共享给我/1、25年年底投标汇总8888/国网安徽电力直属单位片区2025年第七次服务区域联合授权竞争性谈判采购_采购文件包/工作量清单_文件汇总.xlsx"
    search_dir = "/Volumes/zhuxl 共享给我/1、25年年底投标汇总8888/国网安徽电力直属单位片区2025年第七次服务区域联合授权竞争性谈判采购_采购文件包/工程量清单_汇总"
    
    # 创建并运行
    matcher = FileMatcherCopier(excel_path, search_dir)
    matcher.run()


if __name__ == "__main__":
    main()
