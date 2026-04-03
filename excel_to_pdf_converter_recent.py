#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel转PDF工具（带时间过滤版）
只扫描修改时间是2025年12月18日17点30分以后的Excel数据，生成PDF
"""

import os
import glob
import subprocess
import tempfile
import shutil
from pathlib import Path
from datetime import datetime
from PyPDF2 import PdfMerger


def find_soffice():
    """
    查找LibreOffice的soffice可执行文件
    """
    possible_paths = [
        '/opt/homebrew/bin/soffice',
        '/usr/local/bin/soffice',
        '/Applications/LibreOffice.app/Contents/MacOS/soffice',
        shutil.which('soffice'),
        shutil.which('libreoffice'),
    ]
    
    for path in possible_paths:
        if path and os.path.exists(path):
            return path
    
    return None


def find_excel_files_by_time(directory, target_time):
    """
    查找指定目录下的所有Excel文件，且文件名包含“完成报价”，
    且修改时间在 target_time 以后
    
    Args:
        directory: 目标目录路径
        target_time: datetime对象，过滤的起始时间
        
    Returns:
        Excel文件路径列表
    """
    excel_files = []
    # 获取目标时间戳
    target_timestamp = target_time.timestamp()
    
    patterns = ['*.xlsx', '*.xls', '*.xlsm']
    
    for pattern in patterns:
        files = glob.glob(os.path.join(directory, pattern))
        for f in files:
            # 1. 过滤文件名
            if "完成报价" not in os.path.basename(f):
                continue
            
            # 2. 过滤修改时间
            mtime = os.path.getmtime(f)
            if mtime > target_timestamp:
                excel_files.append(f)
                file_time_str = datetime.fromtimestamp(mtime).strftime('%Y-%m-%d %H:%M:%S')
                print(f"  找到匹配文件: {os.path.basename(f)} (修改时间: {file_time_str})")
    
    return sorted(excel_files)


def create_temp_excel_with_sheet(excel_path, sheet_name, temp_dir):
    """
    创建只包含指定sheet的临时Excel文件，并设置页面打印属性
    """
    try:
        import openpyxl
        from openpyxl.worksheet.page import PageMargins
        
        # 打开原始Excel文件
        wb = openpyxl.load_workbook(excel_path)
        
        # 检查sheet是否存在
        if sheet_name not in wb.sheetnames:
            print(f"  警告: 文件 {os.path.basename(excel_path)} 中未找到工作表 '{sheet_name}'")
            wb.close()
            return None
        
        # 创建新工作簿，只包含目标sheet
        new_wb = openpyxl.Workbook()
        new_wb.remove(new_wb.active)  # 删除默认sheet
        
        # 复制目标sheet
        source_sheet = wb[sheet_name]
        target_sheet = new_wb.create_sheet(sheet_name)
        
        # 复制所有单元格（包括样式）
        for row in source_sheet.iter_rows():
            for cell in row:
                target_cell = target_sheet[cell.coordinate]
                target_cell.value = cell.value
                
                # 复制样式
                if cell.has_style:
                    # 使用 copy(obj) 避免 DeprecationWarning
                    from copy import copy
                    target_cell.font = copy(cell.font)
                    target_cell.border = copy(cell.border)
                    target_cell.fill = copy(cell.fill)
                    target_cell.number_format = cell.number_format
                    target_cell.protection = copy(cell.protection)
                    target_cell.alignment = copy(cell.alignment)
        
        # 复制列宽，并稍微增加宽度以防止长数值显示为科学计数法 (######## 或 2E+06)
        for col_letter, col_dim in source_sheet.column_dimensions.items():
            if col_dim.width:
                # 增加 20% 的宽度作为缓冲，确保金额能显示完整
                target_sheet.column_dimensions[col_letter].width = col_dim.width * 1.2
            else:
                # 如果没有定义宽度，设置一个合理的默认宽度
                target_sheet.column_dimensions[col_letter].width = 15
        
        # 复制行高
        for row_num, row_dim in source_sheet.row_dimensions.items():
            target_sheet.row_dimensions[row_num].height = row_dim.height
        
        # 复制合并单元格
        for merged_cell in source_sheet.merged_cells.ranges:
            target_sheet.merge_cells(str(merged_cell))
        
        # 设置页面打印属性 - 关键：让所有列适应在一页宽度内
        target_sheet.page_setup.orientation = 'portrait'  # 竖向
        target_sheet.page_setup.paperSize = 9  # A4纸张
        target_sheet.page_setup.fitToPage = True  # 启用适应页面
        target_sheet.page_setup.fitToWidth = 1  # 所有列适应在1页宽度内
        target_sheet.page_setup.fitToHeight = 0  # 高度不限制（0表示不限制）
        
        # 设置较小的页边距
        target_sheet.page_margins = PageMargins(
            left=0.5, right=0.5, top=0.5, bottom=0.5,
            header=0.3, footer=0.3
        )
        
        # 保存临时文件
        temp_excel_path = os.path.join(temp_dir, f"temp_{os.path.basename(excel_path)}")
        new_wb.save(temp_excel_path)
        
        wb.close()
        new_wb.close()
        
        return temp_excel_path
        
    except Exception as e:
        print(f"  错误: 创建临时Excel文件时出错: {str(e)}")
        return None


def convert_excel_to_pdf_libreoffice(excel_path, output_dir, soffice_path):
    """
    使用LibreOffice将Excel文件转换为PDF
    """
    try:
        cmd = [
            soffice_path,
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
            excel_path
        ]
        
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=60
        )
        
        if result.returncode != 0:
            print(f"  错误: LibreOffice转换失败: {result.stderr}")
            return None
        
        pdf_filename = Path(excel_path).stem + '.pdf'
        pdf_path = os.path.join(output_dir, pdf_filename)
        
        if os.path.exists(pdf_path):
            return pdf_path
        else:
            return None
            
    except Exception as e:
        print(f"  错误: 转换过程出错: {str(e)}")
        return None


def merge_pdfs(pdf_files, output_path):
    """
    合并多个PDF文件为一个
    """
    merger = PdfMerger()
    try:
        for pdf_file in pdf_files:
            if os.path.exists(pdf_file):
                merger.append(pdf_file)
        merger.write(output_path)
        return True
    except Exception as e:
        print(f"合并PDF时出错: {str(e)}")
        return False
    finally:
        merger.close()


def main():
    """主函数"""
    # 查找LibreOffice
    soffice_path = find_soffice()
    if not soffice_path:
        print("错误: 未找到LibreOffice，请先安装LibreOffice")
        return
    
    # 定义过滤时间: 2025年12月18日17点30分
    target_time_limit = datetime(2025, 12, 18, 17, 30, 0)
    print(f"使用LibreOffice: {soffice_path}")
    print(f"过滤条件: 修改时间晚于 {target_time_limit.strftime('%Y-%m-%d %H:%M:%S')}")
    print("-" * 60)
    
    # 目标文件夹路径
    source_directory = input("请输入要处理的文件夹路径: ").strip()
    if not source_directory or not os.path.exists(source_directory):
        print(f"错误: 路径不存在: {source_directory}")
        return
    
    # 查找符合时间条件的Excel文件
    excel_files = find_excel_files_by_time(source_directory, target_time_limit)
    
    if not excel_files:
        print("未找到符合时间条件的Excel文件！")
        return
    
    print(f"\n找到 {len(excel_files)} 个符合条件的Excel文件\n")
    
    # 输出PDF文件路径
    output_directory = source_directory
    timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    final_pdf_name = f"工程量清单_汇总_合并_{timestamp_str}.pdf"
    final_pdf_path = os.path.join(output_directory, final_pdf_name)
    
    # 工作表名称
    sheet_name = "可编辑版"
    
    # 创建临时目录
    temp_dir = tempfile.mkdtemp()
    pdf_output_dir = tempfile.mkdtemp()
    converted_pdfs = []
    
    try:
        for idx, excel_file in enumerate(excel_files, 1):
            print(f"[{idx}/{len(excel_files)}] 处理: {os.path.basename(excel_file)}")
            
            temp_excel = create_temp_excel_with_sheet(excel_file, sheet_name, temp_dir)
            if not temp_excel:
                continue
            
            pdf_path = convert_excel_to_pdf_libreoffice(temp_excel, pdf_output_dir, soffice_path)
            
            if pdf_path:
                final_pdf_name_single = f"{Path(excel_file).stem}_{sheet_name}.pdf"
                final_pdf_path_single = os.path.join(pdf_output_dir, final_pdf_name_single)
                
                if os.path.exists(pdf_path):
                    os.rename(pdf_path, final_pdf_path_single)
                    converted_pdfs.append(final_pdf_path_single)
                    print(f"  ✓ 成功转换")
        
        # 合并PDF
        if converted_pdfs:
            print("\n开始合并PDF文件...")
            if merge_pdfs(converted_pdfs, final_pdf_path):
                print(f"✓ 最终PDF文件: {final_pdf_path}")
                print(f"✓ 成功处理: {len(converted_pdfs)}/{len(excel_files)} 个文件")
        else:
            print("没有成功转换的PDF文件。")
    
    finally:
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir)
        if os.path.exists(pdf_output_dir):
            shutil.rmtree(pdf_output_dir)
        print(f"\n已清理临时文件")


if __name__ == "__main__":
    main()
